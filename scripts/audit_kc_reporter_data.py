#!/usr/bin/env python3
"""Read-only author-workbook extraction; no source imports, simulation or fitting.

Run using the bundled spreadsheet Python runtime (openpyxl, numpy).
"""
from pathlib import Path
from datetime import datetime, timezone
import csv
import hashlib
import json
import subprocess
import xml.etree.ElementTree as ET
import zipfile
import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT / 'data/raw/kc-reporter-data/author-code'
COMMIT = 'f0ee2079dae6761cc2d07f04e96e76e2654b6e3c'
BOOK = REPO / 'data/gamma_mch_responses_manoim_supplement.xlsx'
OUT = ROOT / 'validation/kc-reporter-data-audit.json'
CSV = ROOT / 'validation/kc-reporter-traces.csv'


def pin(path):
    return dict(path=str(path.relative_to(ROOT)), bytes=path.stat().st_size,
                sha256=hashlib.sha256(path.read_bytes()).hexdigest())


def main():
    if OUT.exists() or CSV.exists():
        raise FileExistsError('Preserve the first completed audit')
    def git(*args):
        return subprocess.check_output(['git', '-C', str(REPO), *args], text=True).strip()
    assert git('rev-parse', 'HEAD') == COMMIT
    assert git('status', '--porcelain') == ''
    source_paths = git('ls-files').splitlines()
    sources = [pin(REPO / p) for p in source_paths]
    before = pin(BOOK)
    wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=False)
    assert wb.sheetnames == ['Sheet1']
    ws = wb['Sheet1']
    rows = list(ws.iter_rows())
    assert [[c.value for c in row] for row in rows[:2]] == [
        ['WT', None, 'KD', None], ['Mean', 'SE', 'Mean', 'SE']]
    assert all(c.data_type != 'f' for row in rows for c in row)
    values = np.array([[c.value for c in row] for row in rows[2:]], dtype=float)
    assert values.shape == (499, 4) and np.isfinite(values).all()
    assert (values[:, [1, 3]] > 0).all()
    # Independent OOXML numeric decoding checks every observation and error cell.
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(BOOK) as z:
        xml = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
    xml_values = {c.attrib['r']: float(c.find('s:v', ns).text)
                  for c in xml.findall('.//s:c', ns)
                  if c.attrib.get('t', 'n') == 'n' and c.find('s:v', ns) is not None}
    for i, row in enumerate(values, 3):
        for col, value in zip('ABCD', row):
            assert xml_values[f'{col}{i}'] == value
    wb.close()
    # Reproduce time arithmetic only, without importing or running author code.
    dt = 1 / 30
    t = np.arange(0, len(values) / 30, dt)
    rounded = np.round(np.arange(0, t[-1] + dt, dt), int(np.ceil(-np.log10(dt))))
    if rounded[-1] > t[-1]:
        rounded = rounded[:-1]
    assert len(t) == len(values) == len(rounded)
    stim = (rounded >= 3.6) & (rounded <= 8.6)
    fit = rounded <= 8.6
    with CSV.open('x', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['source_excel_row', 'time_s_from_author_sampling_rate',
                         'rounded_author_time_s', 'author_stimulus_sample',
                         'WT_Mean', 'WT_SE', 'KD_Mean', 'KD_SE'])
        for i, row in enumerate(values):
            writer.writerow([i + 3, t[i], rounded[i], int(stim[i]), *row])
    trace_summary = {}
    for name, col in [('WT', 0), ('KD', 2)]:
        peak = int(np.argmax(values[:, col]))
        trace_summary[name] = dict(
            min_mean=float(values[:, col].min()), max_mean=float(values[peak, col]),
            max_mean_excel_row=peak + 3, max_mean_author_time_s=float(t[peak]),
            min_se=float(values[:, col + 1].min()), max_se=float(values[:, col + 1].max()),
            negative_mean_samples=int((values[:, col] < 0).sum()))
    assert before == pin(BOOK)
    assert git('status', '--porcelain') == ''
    result = dict(
        completed_utc=datetime.now(timezone.utc).isoformat(),
        status='completed_read_only_source_audit',
        repository=dict(url='https://github.com/nawrotlab/KC_KC_lateral_interactions',
                        commit=COMMIT, clean=True, license='GPL-3.0', files=sources),
        workbook=dict(**before, sheet='Sheet1', header_range='A1:D2',
                      observation_range='A3:D501', rows=499, numeric_cells_checked=1996,
                      formula_cells=0, finite=True, positive_SE=True,
                      source_unmodified=True, animal_level_data=False,
                      explicit_timestamps=False, units_in_workbook=False,
                      exact_primary_figure_cohort_match='unresolved', traces=trace_summary),
        author_time_arithmetic=dict(
            sampling_hz=30, declared_onset_s=3.6, declared_duration_s=5,
            first_s=float(t[0]), last_s=float(t[-1]),
            rounded_labels_max_error_s=float(np.max(np.abs(t - rounded))),
            stimulus_sample_indices_inclusive=[int(np.flatnonzero(stim)[0]), int(np.flatnonzero(stim)[-1])],
            stimulus_samples=int(stim.sum()), fit_samples=int(fit.sum()),
            nonzero_drive_euler_intervals=int(stim[:-1].sum()),
            integrated_nonzero_drive_duration_s=float(stim[:-1].sum() * dt),
            derived_time_not_measured=True),
        extraction=pin(CSV), source_script=pin(Path(__file__)),
        neural_simulation_run=False, fit_run=False, runtime_changed=False,
        interpretation='Reporter summary data, not electrical calibration or independent animal trials.')
    with OUT.open('x') as f:
        json.dump(result, f, indent=2)
        f.write('\n')
    print(json.dumps({k: v for k, v in result.items() if k != 'repository'}, indent=2))


if __name__ == '__main__':
    main()
