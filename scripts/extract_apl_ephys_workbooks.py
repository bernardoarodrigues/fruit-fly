"""Read all five APL source workbooks; retain formulas, cached values and identities.

Run with the bundled document Python runtime (openpyxl). Never saves a workbook.
"""
from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
INVENTORY=ROOT/'validation/apl-sk-ephys-inventory.json'
HEADERS=ROOT/'validation/apl-sk-abf-headers.json'
OUT=ROOT/'validation/apl-sk-workbook-join.json'


def color(c):
    return dict(type=c.type,value=c.value,tint=c.tint) if c else None


def main():
    if OUT.exists(): raise FileExistsError('Preserve previous extraction')
    inv=json.loads(INVENTORY.read_text());heads=json.loads(HEADERS.read_text())
    sources=[r for r in inv['entries'] if r['member'].endswith('.xlsx') and '/APL_' in r['member']]
    by_id=defaultdict(list)
    by_filename=defaultdict(list)
    for h in heads['records']:
        by_id[Path(h['candidate_recording_directory']).name].append(h)
        by_filename[Path(h['member']).name.split('_')[0]].append(h)
    books=[];joins=[]
    for source in sources:
        path=ROOT/source['extracted_path']
        assert hashlib.sha256(path.read_bytes()).hexdigest()==source['sha256']
        w=openpyxl.load_workbook(path,data_only=False);cached=openpyxl.load_workbook(path,data_only=True)
        sheets=[]
        for s in w:
            cells=[]
            for row in s:
                for c in row:
                    if c.value is not None or c.comment is not None:
                        cells.append(dict(coordinate=c.coordinate,value=c.value,cached_value=cached[s.title][c.coordinate].value,
                            data_type=c.data_type,number_format=c.number_format,style_id=c.style_id,
                            font_color=color(c.font.color),fill_type=c.fill.patternType,fill_color=color(c.fill.fgColor),
                            strike=c.font.strike,comment=c.comment.text if c.comment else None))
            sheets.append(dict(name=s.title,rows=s.max_row,columns=s.max_column,state=s.sheet_state,
                merged_ranges=[str(x) for x in s.merged_cells.ranges],
                hidden_rows=[k for k,v in s.row_dimensions.items() if v.hidden],
                hidden_columns=[k for k,v in s.column_dimensions.items() if v.hidden],cells=cells))
            if s.title in ('intrinsic','AHP(2nA)'):
                for row_num in range(2,s.max_row+1):
                    cell_id=s.cell(row_num,2).value
                    if not isinstance(cell_id,str) or not cell_id: continue
                    folder_matches=[h for h in by_id.get(cell_id,[]) if h['member'].split('/')[1]==source['member'].split('/')[1]]
                    filename_matches=[h for h in by_filename.get(cell_id,[]) if h['member'].split('/')[1]==source['member'].split('/')[1]]
                    matches=list({h['member']:h for h in folder_matches+filename_matches}.values())
                    genotype=s.cell(row_num,3).value if source['member'].endswith('APL_SKRNAi_final.xlsx') else None
                    state=s.cell(row_num,4 if genotype is not None else 3).value
                    joins.append(dict(workbook=source['member'],sheet=s.title,row=row_num,cell=cell_id,
                        fly=s.cell(row_num,1).value,state=state,genotype=genotype,
                        row_values={s.cell(1,c).coordinate:s.cell(row_num,c).value for c in range(1,s.max_column+1)},
                        directory_id_matches=[h['member'] for h in folder_matches],
                        filename_id_matches=[h['member'] for h in filename_matches],
                        matched_recording_directories=sorted({h['candidate_recording_directory'] for h in matches}),
                        matched_abf_members=[h['member'] for h in matches],
                        join_status='matched_exact_cell_id' if matches else 'unmatched_cell_id'))
        books.append(dict(member=source['member'],path=source['extracted_path'],sha256=source['sha256'],sheets=sheets))
    joined={m for j in joins for m in j['matched_abf_members']}
    result=dict(completed_utc=datetime.now(timezone.utc).isoformat(),openpyxl_version=openpyxl.__version__,
        script_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        inventory_sha256=hashlib.sha256(INVENTORY.read_bytes()).hexdigest(),
        headers_sha256=hashlib.sha256(HEADERS.read_bytes()).hexdigest(),
        workbook_count=len(books),join_row_count=len(joins),join_status_counts=dict(Counter(j['join_status'] for j in joins)),
        unmatched_abf_members=[h['member'] for h in heads['records'] if h['member'] not in joined],
        scope='Read-only source extraction and exact cell-ID joins to directory names or filename prefixes within the same source experiment folder. All formulas/cached values/styles/comments retained. No relabeling, numerical fitting, inferred exclusions, genotype transfer or fuzzy correction.',
        books=books,joins=joins)
    OUT.write_text(json.dumps(result,indent=2,default=str)+'\n')
    print(json.dumps({k:v for k,v in result.items() if k not in ('books','joins','unmatched_abf_members')},indent=2))
    print('Unmatched workbook rows:',[(j['sheet'],j['row'],j['cell']) for j in joins if not j['matched_abf_members']])
    print('Unmatched ABFs:',len(result['unmatched_abf_members']))


if __name__=='__main__':main()
