"""Read original XLSX/PZF bytes; preserve caches, audit simple formulas, no fitting.

Run with the bundled Python runtime. Does not edit/export the source workbook.
Output paths refuse overwrite. PZF handling extracts printable report text only;
it does not claim to parse Prism's binary data table or analysis settings.
"""
from pathlib import Path
from collections import Counter, defaultdict
from decimal import Decimal, localcontext
from xml.etree import ElementTree as ET
from zipfile import ZipFile
import csv
import gzip
import hashlib
import json
import re
import sys

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / 'data/raw/salman-current-clamp'
OUT = ROOT / 'validation'
XLSX = RAW / 'Current Clamp analysis.xlsx'
PZF = RAW / 'ALL CC pre DRUGS.pzf'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def receipt(path):
    data = path.read_bytes()
    return {'path': str(path.relative_to(ROOT)), 'bytes': len(data),
            'sha256': hashlib.sha256(data).hexdigest()}


def json_bytes(obj):
    return (json.dumps(obj, indent=2, allow_nan=False) + '\n').encode()


def main():
    paths = [OUT / 'salman-current-clamp-results.json',
             OUT / 'salman-current-clamp-summary-values.csv',
             OUT / 'salman-current-clamp-prism-text.json',
             RAW / 'xlsx-cells-and-formula-checks.json.gz']
    assert not any(p.exists() for p in paths), 'Refusing to overwrite derived evidence'
    assert receipt(XLSX)['sha256'] == '5e5fd1c48d372369ffdd5f269170d7ac612eedab39bf68233f9f59191cc96fb0'
    assert receipt(PZF)['sha256'] == '55639979681a5a844cff00f24dd386bc99a3fcca5e54f90137741e158ce5d671'
    w = openpyxl.load_workbook(XLSX, data_only=False)
    cache = openpyxl.load_workbook(XLSX, data_only=True)
    with ZipFile(XLSX) as z:
        names = z.namelist()
        core = ET.fromstring(z.read('docProps/core.xml'))
        core_metadata = {node.tag.split('}')[-1]: node.text for node in core}
        original_xml = {}
        # Workbook order is explicitly resolved through relationships, not sheetId.
        relationships = {r.attrib['Id']: r.attrib['Target'] for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
        for node in ET.fromstring(z.read('xl/workbook.xml')).find('s:sheets', NS):
            rid = node.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
            member = 'xl/' + relationships[rid]
            tree = ET.fromstring(z.read(member))
            original_xml[node.attrib['name']] = {c.attrib['r']: {
                'attributes': dict(c.attrib),
                'formula_xml_text': c.findtext('s:f', default=None, namespaces=NS),
                'formula_xml_attributes': dict(c.find('s:f', NS).attrib) if c.find('s:f', NS) is not None else None,
                'cached_xml_text': c.findtext('s:v', default=None, namespaces=NS),
            } for c in tree.findall('.//s:c', NS)}
    raw_cells = []
    formula_checks = []
    sheets = []
    points = []
    max_error = Decimal(0)
    unavailable = []
    errors = []
    jumps = []
    groups = defaultdict(list)
    for s in w:
        group = 'KS' if 'KS' in s.title else 'R32' if 'R32' in s.title else 'ABAF'
        groups[group].append(s.title)
        cells = [c for row in s for c in row if c.value is not None]
        formula_count = 0
        for c in cells:
            cached = cache[s.title][c.coordinate].value
            raw = original_xml[s.title][c.coordinate]
            raw_cells.append({'sheet': s.title, 'address': c.coordinate,
                              'value_or_expanded_formula': c.value, 'cached_value': cached,
                              'data_type': c.data_type, 'number_format': c.number_format, **raw})
            if c.data_type == 'e': errors.append({'sheet': s.title, 'cell': c.coordinate, 'error': c.value})
            if c.data_type != 'f': continue
            formula_count += 1
            # Evaluate only the actually observed, simple local formula vocabulary.
            avg = re.fullmatch(r'=AVERAGE\(([A-Z]+[0-9]+):([A-Z]+[0-9]+)\)', c.value)
            sub = re.fullmatch(r'=([A-Z]+[0-9]+)-([A-Z]+[0-9]+)', c.value)
            with localcontext() as ctx:
                ctx.prec = 50
                if sub:
                    refs = [s[sub[1]], s[sub[2]]]
                    # All subtraction inputs must be original literal numeric cells.
                    assert all(r.data_type != 'f' for r in refs)
                    vals = [Decimal(original_xml[s.title][r.coordinate]['cached_xml_text']) for r in refs]
                    expected = vals[0] - vals[1]
                elif avg:
                    refs = [r for row in s[avg[1]:avg[2]] for r in row]
                    vals = []
                    for r in refs:
                        if r.data_type == 'f':
                            rr = re.fullmatch(r'=([A-Z]+[0-9]+)-([A-Z]+[0-9]+)', r.value)
                            assert rr, (s.title, r.coordinate, r.value)
                            vals.append(Decimal(original_xml[s.title][rr[1]]['cached_xml_text']) -
                                        Decimal(original_xml[s.title][rr[2]]['cached_xml_text']))
                        else:
                            assert isinstance(r.value, (int, float))
                            vals.append(Decimal(original_xml[s.title][r.coordinate]['cached_xml_text']))
                    expected = sum(vals) / len(vals)
                else:
                    raise ValueError(('Unsupported formula; audit stopped', s.title, c.coordinate, c.value))
                if raw['cached_xml_text'] is None:
                    unavailable.append({'sheet': s.title, 'cell': c.coordinate})
                    error = None
                else:
                    error = abs(expected - Decimal(raw['cached_xml_text']))
                    max_error = max(max_error, error)
                formula_checks.append({'sheet': s.title, 'cell': c.coordinate,
                    'expanded_formula': c.value, 'cached_xml_text': raw['cached_xml_text'],
                    'recomputed_from_literal_xml_decimal': str(expected),
                    'absolute_error': str(error) if error is not None else None})
        for ordinal, col in enumerate(range(24, 34), 1):
            c = s.cell(32, col)
            assert c.value == f'=AVERAGE({c.column_letter}1:{c.column_letter}31)'
            points.append({'sheet': s.title, 'group_label_in_sheet_name': group,
                'column_ordinal': ordinal, 'source_cell': c.coordinate, 'formula': c.value,
                'cached_summary_value': cache[s.title][c.coordinate].value,
                'units_in_workbook': '', 'current_pA_in_workbook': '',
                'source_url': 'https://osf.io/sutz6/files/osfstorage/68d98abe626328f871d4e8fd'})
        for block in [(2, 12), (14, 24)]:
            for row in range(2, 32):
                for col in range(*block):
                    a, b = cache[s.title].cell(row - 1, col), cache[s.title].cell(row, col)
                    jumps.append({'sheet': s.title, 'previous_cell': a.coordinate,
                                  'cell': b.coordinate, 'previous': a.value, 'value': b.value,
                                  'absolute_change': abs(b.value - a.value)})
        sheets.append({'sheet': s.title, 'group_label_in_sheet_name': group,
            'declared_dimension': s.calculate_dimension(), 'nonempty_cells': len(cells),
            'formulas': formula_count, 'text_labels': [{'cell': c.coordinate, 'text': c.value} for c in cells if c.data_type == 's'],
            'hidden': s.sheet_state, 'merged_ranges': [str(r) for r in s.merged_cells.ranges],
            'native_table_count': len(s.tables), 'chart_count': len(s._charts),
            'comment_count': sum(c.comment is not None for c in cells),
            'number_formats': dict(Counter(c.number_format for c in cells)),
            'missing_in_raw_blocks_B1_K31_N1_W31': sum(cache[s.title].cell(r, c).value is None for r in range(1, 32) for c in list(range(2, 12)) + list(range(14, 24))),
            'extra_nonempty_after_row32': [c.coordinate for c in cells if c.row > 32],
            'A1_A31_numeric_range': [s['A1'].value, s['A31'].value],
            'M1_M31_numeric_range': [s['M1'].value, s['M31'].value],
            'time_units': None, 'sex': None, 'age': None, 'exact_cell_or_animal_id': None})
    # Preserve an observed anomalous row; do not replace or exclude it.
    dup1 = [cache['Prep R32'].cell(31, c).value for c in range(2, 12)]
    dup2 = [cache['Prep 4 R32'].cell(31, c).value for c in range(2, 12)]
    assert dup1 == dup2
    b = PZF.read_bytes()
    printable = [(m.start(), m.group().decode('ascii')) for m in re.finditer(rb'[\x20-\x7e]{4,}', b)]
    anchors = ['Keystone LNs', 'R32 LNs', 'ABAF LNs']
    report_blocks = []
    for i, (offset, text) in enumerate(printable):
        if text in anchors and i + 16 < len(printable):
            block = printable[i + 1:i + 17]
            if re.fullmatch(r'0\.\d+', block[0][1]) and block[-1][1].startswith('Y = '):
                report_blocks.append({'group': text, 'byte_offset': offset,
                    'fields': dict(zip(['slope', 'intercept', 'x_intercept', 'inverse_slope',
                        'slope_SE', 'intercept_SE', 'slope_CI95', 'intercept_CI95', 'x_intercept_CI95',
                        'R_squared', 'Sy_x', 'F', 'df', 'p', 'significance_label', 'equation'], [v for _, v in block])),
                    'ordered_report_strings': [{'byte_offset': j, 'text': v} for j, v in block]})
    # Instead of inferring binary structure, locate complete human-readable equations.
    equations = [{'byte_offset': o, 'text': t} for o, t in printable if t.startswith('Y = ')]
    significant_report = [{'byte_offset': o, 'prefix': t[:69]} for o, t in printable if t.startswith('Are the slopes equal?')]
    prism = {'source': receipt(PZF), 'method': 'ASCII printable runs >=4 bytes, with byte offsets; no binary Prism table parser',
        'axis_label_strings': [{'byte_offset': o, 'text': t} for o, t in printable if t in ['input pA-', 'Output change mV-']],
        'equations': equations, 'global_test_report_prefixes': significant_report,
        'report_blocks': report_blocks, 'binary_table_or_analysis_settings_verified': False,
        'new_regression_performed': False}
    # No current mapping or fit: means are descriptive over named sheets only.
    means = {g: [{'column_ordinal': j, 'mean_cached_summary': sum(p['cached_summary_value'] for p in points if p['group_label_in_sheet_name'] == g and p['column_ordinal'] == j) / len(names_)} for j in range(1, 11)] for g, names_ in groups.items()}
    with paths[1].open('x', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(points[0])); writer.writeheader(); writer.writerows(points)
    paths[2].write_bytes(json_bytes(prism))
    paths[3].write_bytes(gzip.compress(json_bytes({'cells': raw_cells, 'formula_checks': formula_checks}), mtime=0))
    result = {'source': receipt(XLSX), 'script': receipt(Path(__file__)),
        'runtime': {'python': sys.version, 'openpyxl': openpyxl.__version__},
        'source_core_properties': core_metadata, 'native_package': {'external_links': [n for n in names if 'externalLink' in n], 'table_parts': [n for n in names if '/tables/' in n], 'chart_parts': [n for n in names if '/charts/' in n]},
        'sheet_count': len(sheets), 'group_sheet_names': dict(groups), 'sheets': sheets,
        'formula_audit': {'formulas': len(formula_checks), 'cache_missing': unavailable,
            'literal_error_cells': errors, 'maximum_absolute_cache_error': str(max_error),
            'tolerance': '1e-10 in unlabelled workbook numeric units',
            'all_caches_match_within_tolerance': not unavailable and not errors and max_error < Decimal('1e-10')},
        'descriptive_group_means_by_column': means, 'statistical_independence_of_sheets_established': False,
        'largest_adjacent_changes': sorted(jumps, key=lambda x: -x['absolute_change'])[:20],
        'preserved_anomaly': {'range': 'Prep R32!B31:K31', 'exactly_equals': 'Prep 4 R32!B31:K31', 'values': dup1,
            'included_in_summary': 'Prep R32!X32:AG32', 'corrected_or_excluded': False},
        'current_axis_mapping': {'workbook_headers_available': False, 'assigned_currents': None,
            'published_methods_caption_pA': [-100 + 50*j for j in range(10)],
            'published_figure3D_visually_read_pA': [-200 + 100*j for j in range(10)],
            'discrepancy_resolved': False, 'image_source': receipt(RAW / 'published-figure3.jpg')},
        'published_cohort_match': {'published_cells_per_group': 17, 'workbook_sheets_by_group': {g: len(v) for g, v in groups.items()}, 'exact_match_established': False},
        'derived_files': [receipt(p) for p in paths[1:]], 'source_unchanged': receipt(XLSX)['sha256'] == '5e5fd1c48d372369ffdd5f269170d7ac612eedab39bf68233f9f59191cc96fb0',
        'model_fit': False, 'simulation': False, 'RAR_downloads': 0}
    paths[0].write_bytes(json_bytes(result))
    print(json.dumps({'sheet_count': len(sheets), 'group_counts': {g: len(v) for g, v in groups.items()},
        'formulas': result['formula_audit'], 'prism_equations': len(equations),
        'prism_report_blocks': len(report_blocks), 'result': receipt(paths[0])}, indent=2))


if __name__ == '__main__':
    main()
